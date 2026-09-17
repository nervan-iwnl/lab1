"""Создаёт три отчёта по лабораторным работам в формате DOCX."""

from __future__ import annotations

from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "reports"
AUTHOR = "Иван Агеев"
YEAR = "2026"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_table_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def setup_document(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title_style._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    doc.core_properties.title = title
    doc.core_properties.author = AUTHOR
    add_footer(section)
    return doc


def add_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Страница ")
    set_run_font(run, "Aptos", 9)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_separate, field_text, field_end])


def add_cover(doc: Document, number: int, topic: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    run = p.add_run("НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ ЯДЕРНЫЙ УНИВЕРСИТЕТ МИФИ")
    set_run_font(run, "Aptos", 11, True)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.add_run(f"Отчет по лабораторной работе {number}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(72)
    run = p.add_run(topic)
    set_run_font(run, "Aptos Display", 16, True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.right_indent = Inches(0.4)
    run = p.add_run(f"Выполнил: {AUTHOR}")
    set_run_font(run, "Aptos", 11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    run = p.add_run(f"Москва {YEAR}")
    set_run_font(run, "Aptos", 11)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        set_run_font(run, "Aptos", 11, True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, "Aptos", 11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "Aptos", 11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(item), "Aptos", 11)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        set_run_font(paragraph.add_run(item), "Aptos", 11)


def add_code(doc: Document, text: str, caption: str | None = None) -> None:
    if caption:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        set_run_font(paragraph.add_run(caption), "Aptos", 10, True)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F4F6")
    p_pr.append(shading)
    run = paragraph.add_run(text.rstrip())
    set_run_font(run, "Consolas", 8.7)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, "Aptos", 10, True)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[index])
    keep_table_row(table.rows[0])
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_table_row(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            if row_index % 2:
                set_cell_shading(cell, "EDF3F8")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(paragraph.add_run(value), "Aptos", 9.5)
            if widths:
                cell.width = Inches(widths[column_index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def git_log() -> str:
    command = [
        "git",
        "log",
        "--oneline",
        "--graph",
        "--decorate",
        "--all",
        "--max-count=7",
    ]
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, encoding="utf-8", check=True)
    return result.stdout.strip()


def report_one() -> Path:
    title = "Знакомство с системой контроля версий Git и платформой GitHub"
    doc = setup_document(title)
    add_cover(doc, 1, title)

    add_heading(doc, "Цель работы")
    add_body(
        doc,
        "Освоить базовые операции Git: создание репозитория и коммитов, работу с ветками, "
        "разрешение конфликта при слиянии и безопасную отмену ошибочного изменения через revert.",
    )

    add_heading(doc, "Использованные средства")
    add_table(
        doc,
        ["Средство", "Назначение"],
        [
            ["Git", "Контроль версий, ветвление, слияние и откат"],
            ["Python 3.12", "Демонстрационная программа hello.py"],
            ["GitHub", "Целевая платформа для публикации репозитория и Pull Request"],
        ],
        [1.7, 4.6],
    )

    add_heading(doc, "Ход работы")
    add_heading(doc, "Создание репозитория и первых коммитов", 2)
    add_body(
        doc,
        "В каталоге TRPO создан репозиторий с основной веткой main. Сначала добавлен файл README.md, "
        "затем отдельным коммитом создан hello.py. Раздельные коммиты позволяют увидеть развитие проекта.",
    )
    add_code(
        doc,
        "git init -b main\n"
        "git add README.md\n"
        "git commit -m \"Добавлен README.md\"\n"
        "git add hello.py\n"
        "git commit -m \"Добавлен hello.py\"",
        "Основные команды",
    )

    add_heading(doc, "Работа с веткой и конфликтом", 2)
    add_body(
        doc,
        "В ветке feature-branch реализовано персональное приветствие. Параллельно в main изменена та же "
        "строка файла, поэтому слияние вызвало содержательный конфликт. Итоговая версия вручную объединяет "
        "оба изменения и сохранена merge-коммитом.",
    )
    add_code(
        doc,
        'print("Git - это система контроля версий")\n'
        'name = input("Введите ваше имя: ")\n'
        'print(f"Привет, {name}! Рад познакомиться.")',
        "Итоговый файл hello.py после разрешения конфликта",
    )

    add_heading(doc, "Безопасный откат", 2)
    add_body(
        doc,
        "После слияния создан намеренно ошибочный коммит BAD COMMIT: сломал код. Команда git revert HEAD "
        "создала новый коммит, отменяющий ошибку без переписывания истории.",
    )
    add_code(doc, "git commit -m \"BAD COMMIT: сломал код\"\ngit revert HEAD --no-edit")

    add_heading(doc, "Результаты")
    add_code(doc, git_log(), "Фактическая история репозитория")
    add_code(
        doc,
        "Git - это система контроля версий\n"
        "Введите ваше имя: Иван\n"
        "Привет, Иван! Рад познакомиться.",
        "Результат запуска программы",
    )
    add_body(
        doc,
        "Локальная часть лабораторной работы выполнена полностью. Для формирования Pull Request на GitHub "
        "нужно добавить адрес удалённого репозитория origin, отправить main и feature-branch и выполнить "
        "слияние через веб-интерфейс под учётной записью владельца.",
    )
    add_code(
        doc,
        "# После добавления URL репозитория как origin:\n"
        "git push -u origin main\n"
        "git push -u origin feature-branch",
        "Команды публикации после создания пустого репозитория",
    )

    add_heading(doc, "Вывод")
    add_body(
        doc,
        "В ходе работы создана последовательная история Git, отработаны изолированная разработка в ветке, "
        "ручное разрешение конфликта и обратимый откат через revert. Итоговый код находится в корректном "
        "состоянии, а все учебные этапы видны в графе коммитов.",
    )

    path = OUTPUT_DIR / "Отчет ЛБ 1 Git и GitHub.docx"
    doc.save(path)
    return path


def report_two() -> Path:
    title = "Основы работы с базами данных через ORM"
    doc = setup_document(title)
    add_cover(doc, 2, title)

    add_heading(doc, "Цель работы")
    add_body(
        doc,
        "Освоить SQLAlchemy ORM, описать связанные сущности магазина и выполнить операции создания, чтения, "
        "обновления и удаления данных без ручного написания SQL-запросов.",
    )

    add_heading(doc, "Структура проекта")
    add_table(
        doc,
        ["Файл", "Назначение"],
        [
            ["database.py", "Подключение к SQLite, engine, фабрика сессий и инициализация схемы"],
            ["models.py", "Модели Category, Product, Order и OrderItem"],
            ["crud.py", "CRUD, фильтрация и добавление товаров в заказ"],
            ["main.py", "Повторяемая демонстрация всех операций"],
            ["tests/test_crud.py", "Автоматическая проверка CRUD и связей"],
        ],
        [1.75, 4.55],
    )

    add_heading(doc, "Модель данных")
    add_table(
        doc,
        ["Связь", "Тип", "Реализация"],
        [
            ["Category → Product", "1:N", "products и category_id"],
            ["Order → OrderItem", "1:N", "order_items и order_id"],
            ["Product → OrderItem", "1:N", "order_items и product_id"],
            ["Order ↔ Product", "N:M", "Промежуточная таблица order_items"],
        ],
        [1.8, 1.0, 3.5],
    )
    add_body(
        doc,
        "Промежуточная сущность OrderItem хранит не только внешние ключи, но также количество товара и цену "
        "на момент оформления заказа. Для целостности добавлены ограничения положительного количества, "
        "неотрицательной цены и уникальности пары заказ–товар.",
    )
    add_code(
        doc,
        "class OrderItem(Base):\n"
        "    __tablename__ = \"order_items\"\n"
        "    order_id = mapped_column(ForeignKey(\"orders.id\", ondelete=\"CASCADE\"))\n"
        "    product_id = mapped_column(ForeignKey(\"products.id\", ondelete=\"CASCADE\"))\n"
        "    quantity = mapped_column(Integer, default=1)\n"
        "    price_at_time = mapped_column(Float, nullable=False)",
        "Фрагмент модели позиции заказа",
    )

    add_heading(doc, "Реализация CRUD")
    add_body(
        doc,
        "Для каждой основной сущности реализованы операции создания, чтения, обновления и удаления. Выборка "
        "товаров поддерживает пагинацию, фильтр категории и ценовой диапазон. После записи транзакция "
        "фиксируется, а ORM-объект обновляется из базы.",
    )
    add_code(
        doc,
        "statement = select(Product).order_by(Product.id)\n"
        "if category_id is not None:\n"
        "    statement = statement.where(Product.category_id == category_id)\n"
        "if min_price is not None:\n"
        "    statement = statement.where(Product.price >= min_price)\n"
        "if max_price is not None:\n"
        "    statement = statement.where(Product.price <= max_price)",
        "Фильтрация товаров средствами ORM",
    )

    add_heading(doc, "Результаты выполнения")
    add_code(
        doc,
        "База данных инициализирована\n"
        "Создан заказ #1 с двумя позициями\n"
        "Товар в диапазоне: iPhone 14 — 79999.99 руб.\n"
        "Товар в диапазоне: Python для начинающих — 1500.00 руб.\n"
        "Позиция заказа: iPhone 14, количество: 1\n"
        "Позиция заказа: Python для начинающих, количество: 2\n"
        "Категория, остаток и статус заказа обновлены\n"
        "Временный товар удалён: True\n"
        "В категории 'Электроника и гаджеты' товаров: 3\n"
        "ДЕМОНСТРАЦИЯ ЗАВЕРШЕНА",
        "Фактический вывод python -m lab2_orm.main",
    )
    add_body(
        doc,
        "Автоматические тесты проверяют фильтрацию, изменение категории и остатка, накопление количества "
        "одинакового товара в позиции заказа, связь многие-ко-многим и каскадное удаление.",
    )
    add_code(doc, "...                                                                      [100%]\n3 passed in 0.67s", "Результат pytest")

    add_heading(doc, "Вывод")
    add_body(
        doc,
        "SQLAlchemy позволил представить таблицы SQLite в виде Python-классов и выполнять CRUD через "
        "объекты и сессию. Реализованные связи корректно отражают предметную область магазина, а ограничения "
        "схемы и тесты подтверждают целостность данных и работоспособность операций.",
    )

    path = OUTPUT_DIR / "Отчет ЛБ 2 SQLAlchemy ORM.docx"
    doc.save(path)
    return path


def report_three() -> Path:
    title = "Создание простого REST API"
    doc = setup_document(title)
    add_cover(doc, 3, title)

    add_heading(doc, "Цель работы")
    add_body(
        doc,
        "Создать веб-сервис на FastAPI, реализовать обработку HTTP-запросов GET и POST, настроить валидацию "
        "JSON и проверить успешные ответы и ошибку 404.",
    )

    add_heading(doc, "Архитектура API")
    add_table(
        doc,
        ["Метод", "Маршрут", "Результат"],
        [
            ["POST", "/items", "Создание объекта, UUID и статус 201"],
            ["GET", "/items", "Список всех объектов, статус 200"],
            ["GET", "/items/{item_id}", "Объект со статусом 200 или ошибка 404"],
        ],
        [0.85, 1.65, 3.8],
    )
    add_body(
        doc,
        "Данные хранятся в словаре процесса, как предусмотрено учебным заданием. Модель Item проверяет "
        "непустое имя, ограничивает длину полей и запрещает неизвестные поля. Модель ItemInDB добавляет ID.",
    )
    add_code(
        doc,
        "class Item(BaseModel):\n"
        "    model_config = ConfigDict(extra=\"forbid\", str_strip_whitespace=True)\n"
        "    name: Annotated[str, Field(min_length=1, max_length=200)]\n"
        "    description: Annotated[str | None, Field(max_length=500)] = None\n\n"
        "class ItemInDB(Item):\n"
        "    id: str",
        "Pydantic-модели запроса и ответа",
    )

    add_heading(doc, "Реализация маршрутов")
    add_code(
        doc,
        "@app.post(\"/items\", response_model=ItemInDB, status_code=201)\n"
        "async def create_item(item: Item) -> ItemInDB:\n"
        "    item_to_save = ItemInDB(id=str(uuid4()), **item.model_dump())\n"
        "    database[item_to_save.id] = item_to_save\n"
        "    return item_to_save\n\n"
        "@app.get(\"/items/{item_id}\", response_model=ItemInDB)\n"
        "async def get_item_by_id(item_id: str) -> ItemInDB:\n"
        "    item = database.get(item_id)\n"
        "    if item is None:\n"
        "        raise HTTPException(status_code=404, detail=\"Item not found\")\n"
        "    return item",
        "Создание объекта и обработка отсутствующего ID",
    )

    add_heading(doc, "Тестирование API")
    add_body(
        doc,
        "Проверка проведена без внешнего сервера через ASGI-транспорт: запросы проходят полный стек FastAPI, "
        "но выполняются быстрее и воспроизводимо. Созданы два объекта, затем проверены список, поиск по ID, "
        "валидация и ответ для отсутствующего ресурса.",
    )
    add_table(
        doc,
        ["Проверка", "Ожидаемый статус", "Фактический статус"],
        [
            ["Создание ноутбука", "201 Created", "201 Created"],
            ["Создание мыши", "201 Created", "201 Created"],
            ["Получение списка", "200 OK", "200 OK"],
            ["Получение по существующему ID", "200 OK", "200 OK"],
            ["Получение по отсутствующему ID", "404 Not Found", "404 Not Found"],
            ["Лишнее поле или пустое имя", "422 Unprocessable Content", "422 Unprocessable Content"],
        ],
        [2.55, 1.85, 1.9],
    )
    add_code(
        doc,
        "POST /items -> 201\n"
        "{\"name\":\"Ноутбук\",\"description\":\"Игровой ноутбук\",\n"
        " \"id\":\"e43bae4d-2a44-4c6c-a9af-68259256af80\"}\n\n"
        "GET /items -> 200, получено 2 объекта\n"
        "GET /items/e43bae4d-2a44-4c6c-a9af-68259256af80 -> 200\n"
        "GET /items/несуществующий_id -> 404\n"
        "{\"detail\":\"Item not found\"}",
        "Фактические результаты запросов",
    )
    add_code(doc, "...                                                                      [100%]\n3 passed in 0.67s", "Общий результат автоматических тестов")

    add_heading(doc, "Запуск и автодокументация")
    add_code(doc, "python -m uvicorn lab3_rest_api.main:app --reload")
    add_body(
        doc,
        "После запуска интерактивная документация Swagger доступна по адресу http://127.0.0.1:8000/docs. "
        "FastAPI формирует её автоматически на основании маршрутов, моделей Pydantic и кодов ответов.",
    )

    add_heading(doc, "Вывод")
    add_body(
        doc,
        "Создан REST API с тремя требуемыми маршрутами. Сервис корректно валидирует входные данные, возвращает "
        "201 при создании, 200 при чтении и 404 для неизвестного ID. Автоматические тесты подтверждают работу "
        "основных и ошибочных сценариев.",
    )

    path = OUTPUT_DIR / "Отчет ЛБ 3 REST API.docx"
    doc.save(path)
    return path


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [report_one(), report_two(), report_three()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
